#!/usr/bin/env python3
"""
Script to generate a Word document with complete project instructions.
Requires: python-docx library
Install with: pip install python-docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

# Create a new Document
doc = Document()

# Add title
title = doc.add_heading('Cyber Awareness Game', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add subtitle
subtitle = doc.add_paragraph('Complete Setup and Running Instructions')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_format = subtitle.runs[0]
subtitle_format.font.size = Pt(14)
subtitle_format.font.italic = True

# Add date
date_para = doc.add_paragraph(f'Generated: {datetime.now().strftime("%B %d, %Y")}')
date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_para.runs[0].font.size = Pt(10)
date_para.runs[0].font.color.rgb = RGBColor(128, 128, 128)

doc.add_paragraph()  # Spacing

# Prerequisites Section
doc.add_heading('Prerequisites', level=2)
prereq_text = doc.add_paragraph(
    'Before starting, ensure you have installed the following:'
)

# Prerequisites list
prereqs = [
    'Node.js (v14 or higher) - Download from https://nodejs.org/',
    'MongoDB (v4.4 or higher) - Download from https://www.mongodb.com/try/download/community',
    'npm (comes with Node.js)',
    'Python 3.7+ (for generating this documentation)'
]

for prereq in prereqs:
    doc.add_paragraph(prereq, style='List Bullet')

doc.add_paragraph()  # Spacing

# Setup Instructions
doc.add_heading('Setup Instructions', level=2)

# Step 1
doc.add_heading('Step 1: Start MongoDB', level=3)
step1_text = [
    'Open Command Prompt or PowerShell',
    'Start MongoDB by running:',
]
for text in step1_text:
    doc.add_paragraph(text)

code_para = doc.add_paragraph('mongod')
code_para_format = code_para.runs[0]
code_para_format.font.name = 'Courier New'
code_para_format.font.size = Pt(10)
code_para_format.bold = True
code_para.paragraph_format.left_indent = Inches(0.5)

doc.add_paragraph(
    'This starts the MongoDB server on localhost:27017. Keep this terminal window open.',
    style='List Bullet'
)

doc.add_paragraph()  # Spacing

# Step 2
doc.add_heading('Step 2: Install Dependencies', level=3)
doc.add_paragraph(
    'Open a new terminal/PowerShell window in the project directory (c:\\Users\\rajas\\Downloads\\cyber-awareness-game):'
)

doc.add_heading('Install Frontend Dependencies:', level=4)
code_para = doc.add_paragraph('npm install')
code_para_format = code_para.runs[0]
code_para_format.font.name = 'Courier New'
code_para_format.font.size = Pt(10)
code_para_format.bold = True
code_para.paragraph_format.left_indent = Inches(0.5)

doc.add_heading('Install Backend Dependencies:', level=4)
install_steps = [
    'cd server',
    'npm install',
    'cd ..'
]
for step in install_steps:
    code_para = doc.add_paragraph(step)
    code_para_format = code_para.runs[0]
    code_para_format.font.name = 'Courier New'
    code_para_format.font.size = Pt(10)
    code_para_format.bold = True
    code_para.paragraph_format.left_indent = Inches(0.5)

doc.add_paragraph()  # Spacing

# Step 3
doc.add_heading('Step 3: Start the Backend Server', level=3)
doc.add_paragraph('In the project root, start the backend:')

backend_steps = [
    'cd server',
    'npm start'
]
for step in backend_steps:
    code_para = doc.add_paragraph(step)
    code_para_format = code_para.runs[0]
    code_para_format.font.name = 'Courier New'
    code_para_format.font.size = Pt(10)
    code_para_format.bold = True
    code_para.paragraph_format.left_indent = Inches(0.5)

doc.add_paragraph('Or if npm start is not configured, run:', style='List Bullet')

code_para = doc.add_paragraph('node server.js')
code_para_format = code_para.runs[0]
code_para_format.font.name = 'Courier New'
code_para_format.font.size = Pt(10)
code_para_format.bold = True
code_para.paragraph_format.left_indent = Inches(0.75)

doc.add_paragraph(
    'The backend server will start on http://localhost:5000 (or check the server.js for the configured port). Keep this terminal open.',
    style='List Bullet'
)

doc.add_paragraph()  # Spacing

# Step 4
doc.add_heading('Step 4: Start the Frontend Development Server', level=3)
doc.add_paragraph('Open another terminal/PowerShell window in the project root:')

code_para = doc.add_paragraph('npm run dev')
code_para_format = code_para.runs[0]
code_para_format.font.name = 'Courier New'
code_para_format.font.size = Pt(10)
code_para_format.bold = True
code_para.paragraph_format.left_indent = Inches(0.5)

doc.add_paragraph(
    'This will start the Vite development server on http://localhost:5173 (or display the actual port in the terminal).',
    style='List Bullet'
)

doc.add_paragraph()  # Spacing

# Step 5
doc.add_heading('Step 5: Access the Application', level=3)
access_steps = [
    'Open your browser',
    'Navigate to http://localhost:5173',
    'The application should now load and be fully functional'
]
for i, step in enumerate(access_steps, 1):
    doc.add_paragraph(step, style='List Number')

doc.add_paragraph()  # Spacing

# Summary
doc.add_heading('Summary of Running Processes', level=2)
doc.add_paragraph(
    'You should have 3 terminal windows open simultaneously:'
)

processes = [
    'Terminal 1: MongoDB (mongod)',
    'Terminal 2: Backend Server (node server.js or from server/ folder)',
    'Terminal 3: Frontend Development Server (npm run dev from root)'
]
for process in processes:
    doc.add_paragraph(process, style='List Number')

doc.add_paragraph()  # Spacing

# Available Commands
doc.add_heading('Available Commands', level=2)

doc.add_heading('Frontend Commands (from root):', level=3)
frontend_commands = [
    'npm run dev - Start development server',
    'npm run build - Build for production',
    'npm run preview - Preview production build',
    'npm run lint - Run ESLint'
]
for cmd in frontend_commands:
    doc.add_paragraph(cmd, style='List Bullet')

doc.add_heading('Backend Commands:', level=3)
doc.add_paragraph('Start with: node server.js from the server/ folder', style='List Bullet')

doc.add_paragraph()  # Spacing

# Troubleshooting
doc.add_heading('Troubleshooting', level=2)

troubleshoots = [
    '"MongoDB connection failed" - Ensure MongoDB is running with mongod',
    '"Port already in use" - Change the port in server.js or stop conflicting applications',
    '"Module not found" - Run npm install in both root and server/ directories',
    '"Frontend not loading" - Check that both backend and frontend servers are running'
]
for troubleshoot in troubleshoots:
    doc.add_paragraph(troubleshoot, style='List Bullet')

doc.add_paragraph()  # Spacing

# Project Structure
doc.add_heading('Project Structure', level=2)

structure_info = doc.add_paragraph(
    'The project is organized as follows:'
)

structure_items = [
    'Root: Frontend application (React + Vite)',
    'server/: Backend application (Express + MongoDB)',
    'src/: Frontend source code (React components, pages, utilities)',
    'public/: Static assets (sounds, images, etc.)',
    'src/data/: Game data and story definitions',
    'src/pages/: Page components for routing'
]
for item in structure_items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph()  # Spacing

# Footer
footer_para = doc.add_paragraph()
footer_para.paragraph_format.space_before = Pt(12)
footer_run = footer_para.add_run('For more information, refer to the README.md file in the project root.')
footer_run.font.size = Pt(10)
footer_run.font.italic = True
footer_run.font.color.rgb = RGBColor(128, 128, 128)

# Save the document
output_path = 'Instructions.docx'
doc.save(output_path)
print(f'✓ Document created successfully: {output_path}')
print(f'✓ Location: {output_path}')
